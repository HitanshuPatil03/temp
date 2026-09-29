"""End-to-end ingestion: a PDF goes in, evidence-backed facts come out.

This is the file that proves the product's claim rather than its parts. A
synthetic CIL-style production table is built with real ruling lines, run through
every stage in order, and then checked for the things that actually matter:

- the figures are right, in canonical units, with the *printed* value kept;
- every fact carries a page, a table, a cell and a verbatim snippet;
- the total row is **not** extracted (it would double-count);
- a second run produces the same rows, not duplicates — the property that makes
  killing a worker mid-ingest safe;
- the state machine refuses to skip a stage.

The stages are invoked directly rather than through a worker process, because a
worker commits its own transactions and the test fixture holds one open. What is
exercised is the same code the worker calls, in the same order, with the same
store.
"""

from __future__ import annotations

import pymupdf
import pytest
import sqlalchemy as sa

from mrip.blobs import FilesystemBlobStore
from mrip.db.tables import jobs
from mrip.ingest import pipeline
from mrip.ingest.lifecycle import IllegalTransitionError, require_transition
from mrip.jobs.queue import JobQueue
from mrip.jobs.registry import JobContext
from mrip.schemas import DocumentState, FactStatus

WORKER = "test-worker:1"

#: The table this suite ingests. Deliberately shaped like a real CIL annual
#: report page: the metric is in the caption, the unit is in the caption, the
#: rows are subsidiaries and the columns are fiscal years.
CAPTION = "Coal production by subsidiary"
UNIT_NOTE = "(Figures in Million Tonnes)"
TABLE = [
    ["Subsidiary", "FY2023-24", "FY2024-25"],
    ["SECL", "167.00", "193.00"],
    ["MCL", "193.00", "210.50"],
    ["Total", "360.00", "403.50"],
]


def build_report_pdf(path) -> None:
    """A one-page PDF with a ruled table, written with PyMuPDF.

    Ruling lines matter: with them pdfplumber finds a *lattice* table, whose
    structure is read from the document rather than inferred from whitespace.
    That is the higher-confidence path, and the one a real report uses.
    """
    document = pymupdf.open()
    page = document.new_page()

    page.insert_text((72, 72), CAPTION, fontsize=13)
    page.insert_text((72, 90), UNIT_NOTE, fontsize=9)

    left, top = 72, 110
    row_height, col_width = 22, 110
    for row_index, row in enumerate(TABLE):
        for col_index, value in enumerate(row):
            page.insert_text(
                (left + col_index * col_width + 6, top + row_index * row_height + 15),
                value,
                fontsize=10,
            )

    rows, cols = len(TABLE), len(TABLE[0])
    for row_index in range(rows + 1):
        y = top + row_index * row_height
        page.draw_line((left, y), (left + cols * col_width, y))
    for col_index in range(cols + 1):
        x = left + col_index * col_width
        page.draw_line((x, top), (x, top + rows * row_height))

    document.save(path)
    document.close()


@pytest.fixture
def blobs(tmp_path, monkeypatch) -> FilesystemBlobStore:
    """A blob store under the test's temp directory, used by the pipeline."""
    store = FilesystemBlobStore(tmp_path / "blobs")
    monkeypatch.setattr(pipeline, "get_blob_store", lambda: store)
    return store


@pytest.fixture
def ingested(tmp_path, store, blobs, make_document):
    """A registered document whose bytes are in the blob store, ready to run."""
    path = tmp_path / "cil-production.pdf"
    build_report_pdf(path)
    reference = blobs.put_file(path)

    document = store.register_document(
        make_document(
            "cil-production.pdf",
            document_id="doc_pipeline",
            content_hash=reference.content_hash,
            size_bytes=reference.size_bytes,
            page_count=1,
            publisher_entity_id="cil",
            fiscal_year=None,
            blob_key=reference.content_hash,
        )
    )
    return document


def run_stage(store, document_id: str, kind: str) -> None:
    """Invoke one stage the way the worker does."""
    from mrip.jobs.registry import get_handler

    job = JobQueue(store.connection).enqueue(kind, {"document_id": document_id}).job
    get_handler(kind)(
        JobContext(job=job, store=store, worker_id=WORKER, lease_seconds=60)
    )


def run_pipeline(store, document_id: str) -> None:
    for kind in (
        "document.classify",
        "document.digitize",
        "document.extract",
        "document.normalize",
        "document.validate",
        "document.index",
    ):
        run_stage(store, document_id, kind)


# ------------------------------------------------------------------ the chain


def test_a_pdf_becomes_evidence_backed_facts(store, ingested, scope):
    """The whole claim, in one test."""
    run_pipeline(store, ingested.document_id)

    document = store.get_document(ingested.document_id, scope)
    assert document.state is DocumentState.READY

    facts = store.query_facts(scope, include_inactive=True, limit=100)
    by_key = {(fact.entity_id, fact.fiscal_year): fact for fact in facts}

    # Two subsidiaries × two fiscal years. The total row is excluded.
    assert set(by_key) == {
        ("secl", "FY2023-24"),
        ("secl", "FY2024-25"),
        ("mcl", "FY2023-24"),
        ("mcl", "FY2024-25"),
    }

    secl = by_key[("secl", "FY2024-25")]
    assert secl.metric == "coal_production"
    # 193 million tonnes, canonicalised to tonnes…
    assert secl.value == pytest.approx(193_000_000.0)
    assert secl.unit == "t"
    # …with what the page actually printed kept alongside it.
    assert secl.raw_value == pytest.approx(193.0)
    assert "193.00" in (secl.evidence.snippet or "")


def test_every_fact_can_be_traced_back_to_a_cell(store, ingested, scope):
    """The central invariant: no figure without a receipt."""
    run_pipeline(store, ingested.document_id)

    for fact in store.query_facts(scope, include_inactive=True, limit=100):
        evidence = fact.evidence
        assert evidence.document_id == ingested.document_id
        assert evidence.page == 1
        assert evidence.table_id, "a table cell must name its table"
        assert evidence.cell_ref, "and its position inside it"
        assert evidence.bbox is not None, "and where to draw the highlight"
        assert evidence.snippet, "and quote the source"
        assert "doc:" in evidence.locator


def test_a_total_row_is_not_extracted(store, ingested, scope):
    """Storing a total beside its parts double-counts every aggregate."""
    run_pipeline(store, ingested.document_id)

    values = {fact.raw_value for fact in store.query_facts(scope, limit=100)}
    assert 360.00 not in values
    assert 403.50 not in values


def test_confidence_stays_decomposed_and_names_its_weakest_stage(store, ingested, scope):
    """A lattice table has no OCR stage, so `ocr` stays null rather than 1.0."""
    run_pipeline(store, ingested.document_id)

    fact = store.query_facts(scope, limit=1)[0]
    assert fact.confidence.ocr is None, "no OCR ran; that is not the same as perfect"
    assert fact.confidence.answer is None, "no model was involved in a table cell"
    assert fact.confidence.parse is not None
    assert fact.confidence.limiting == fact.confidence.parse


def test_the_document_learns_its_own_fiscal_years_and_entities(store, ingested, scope):
    """The index stage derives facets from the facts, not from the filename."""
    run_pipeline(store, ingested.document_id)

    facets = store.facts.document_facets(ingested.document_id, 1)
    assert facets["entities"] == ["mcl", "secl"]
    assert facets["fiscal_years"] == ["FY2023-24", "FY2024-25"]
    assert facets["metrics"] == ["coal_production"]


# ------------------------------------------------------------------ idempotency


def test_re_running_a_stage_replaces_its_output_rather_than_doubling_it(
    store, ingested, scope
):
    """The property that makes `kill -9` mid-ingest safe."""
    run_stage(store, ingested.document_id, "document.classify")
    run_stage(store, ingested.document_id, "document.digitize")
    first = store.evidence.for_document(ingested.document_id, 1)

    # As if the worker died after writing evidence but before committing its
    # state change, and the job was reclaimed.
    store.documents.set_state(ingested.document_id, DocumentState.CLASSIFIED)
    run_stage(store, ingested.document_id, "document.digitize")
    second = store.evidence.for_document(ingested.document_id, 1)

    assert len(second) == len(first), "a retry must not duplicate evidence"
    assert [row["text"] for row in second] == [row["text"] for row in first]


def test_re_extracting_replaces_facts_rather_than_duplicating_them(
    store, ingested, scope
):
    run_pipeline(store, ingested.document_id)
    before = len(store.query_facts(scope, include_inactive=True, limit=100))

    store.documents.set_state(ingested.document_id, DocumentState.DIGITIZED)
    run_stage(store, ingested.document_id, "document.extract")

    after = len(store.query_facts(scope, include_inactive=True, limit=100))
    assert after == before


# --------------------------------------------------------------- state machine


def test_the_pipeline_advances_one_state_at_a_time(store, ingested, scope):
    expected = [
        ("document.classify", DocumentState.CLASSIFIED),
        ("document.digitize", DocumentState.DIGITIZED),
        ("document.extract", DocumentState.EXTRACTED),
        ("document.normalize", DocumentState.NORMALIZED),
        ("document.validate", DocumentState.VALIDATED),
    ]
    for kind, state in expected:
        run_stage(store, ingested.document_id, kind)
        assert store.get_document(ingested.document_id, scope).state is state


def test_a_stage_cannot_be_skipped(store, ingested, scope):
    """Ordering alone would allow `received → ready`; the transition table does not."""
    with pytest.raises(IllegalTransitionError, match="cannot move to"):
        require_transition(DocumentState.RECEIVED, DocumentState.READY)

    # And the handler refuses too, rather than writing a state its work does not
    # support.
    with pytest.raises(IllegalTransitionError):
        run_stage(store, ingested.document_id, "document.extract")


def test_each_stage_enqueues_the_next_one(store, ingested):
    """The chain is driven by the queue, in the same transaction as the work.

    Not "a job exists somewhere": the *successor* of the stage that just ran has
    to be pending, or the document stops silently after one step.
    """
    queue = JobQueue(store.connection)
    run_stage(store, ingested.document_id, "document.classify")

    kinds = {
        row["kind"]
        for row in store.connection.execute(sa.select(jobs)).mappings()
        if row["state"] == "pending"
    }
    assert "document.digitize" in kinds
    assert queue.depth()["pending"] >= 1


def test_a_missing_document_fails_permanently_rather_than_retrying(store):
    """Three attempts at a document that was deleted is three wasted leases."""
    from mrip.jobs.queue import PermanentJobError

    with pytest.raises(PermanentJobError, match="does not exist"):
        run_stage(store, "doc_gone", "document.classify")


def test_a_document_with_no_stored_bytes_fails_permanently(store, make_document):
    """A document row that survived a restore its blob volume did not."""
    from mrip.jobs.queue import PermanentJobError

    document = store.register_document(
        make_document("orphan.pdf", document_id="doc_orphan", blob_key=None)
    )
    with pytest.raises(PermanentJobError, match="no blob key"):
        run_stage(store, document.document_id, "document.classify")


# ----------------------------------------------------------------- validation


def test_a_figure_that_exceeds_national_production_is_flagged(
    store, ingested, make_fact, scope
):
    """The unit-misread check, on a figure no mine could produce."""
    from mrip.validate.rules import run_rules

    store.insert_facts(
        [
            make_fact(
                document_id=ingested.document_id,
                value=5.0e9,
                raw_value=5000.0,
                raw_unit="MT",
            )
        ]
    )
    findings = run_rules(store, ingested)

    assert any(finding.rule == "implausible_magnitude" for finding in findings)
    flagged = store.query_facts(scope, status=FactStatus.NEEDS_REVIEW, limit=10)
    assert flagged, "the fact was routed to review, not deleted or corrected"


def test_a_negative_production_figure_is_flagged(store, ingested, make_fact, scope):
    from mrip.validate.rules import run_rules

    store.insert_facts(
        [make_fact(document_id=ingested.document_id, value=-193.0e6, raw_value=-193.0)]
    )
    findings = run_rules(store, ingested)

    assert any(finding.rule == "negative_value" for finding in findings)
    assert "not physically possible" in findings[0].message


# --------------------------------------------------------------- other classes


def build_docx(path) -> None:
    """A Word file shaped like a monthly note: a heading and a small table."""
    import docx

    document = docx.Document()
    document.add_paragraph("Coal production by subsidiary")
    document.add_paragraph("(Figures in Million Tonnes)")

    table = document.add_table(rows=len(TABLE), cols=len(TABLE[0]))
    for row_index, row in enumerate(TABLE):
        for col_index, value in enumerate(row):
            table.cell(row_index, col_index).text = value

    document.save(path)


def test_a_word_document_yields_facts_with_paragraph_and_cell_locators(
    tmp_path, store, blobs, make_document, scope
):
    """Word files are a real CIL input — a note is written in Word long before it
    becomes a PDF — and the PS names "digital documents" explicitly."""
    from mrip.schemas import DocumentClass

    path = tmp_path / "monthly-note.docx"
    build_docx(path)
    reference = blobs.put_file(path)

    document = store.register_document(
        make_document(
            "monthly-note.docx",
            document_id="doc_word",
            content_hash=reference.content_hash,
            size_bytes=reference.size_bytes,
            doc_class=DocumentClass.DOCX,
            page_count=None,
            publisher_entity_id="cil",
            fiscal_year=None,
            blob_key=reference.content_hash,
        )
    )

    run_pipeline(store, document.document_id)

    assert store.get_document(document.document_id, scope).state is DocumentState.READY

    facts = store.query_facts(scope, include_inactive=True, limit=100)
    assert {(fact.entity_id, fact.fiscal_year) for fact in facts} == {
        ("secl", "FY2023-24"),
        ("secl", "FY2024-25"),
        ("mcl", "FY2023-24"),
        ("mcl", "FY2024-25"),
    }

    # A .docx has no pages — Word computes pagination at render time — so the
    # locator leans on the table cell rather than inventing a page number.
    for fact in facts:
        assert fact.evidence.page is None, "a page number here would be invented"
        assert fact.evidence.cell_ref


def test_an_image_without_the_ocr_extra_fails_with_an_actionable_reason(
    tmp_path, store, blobs, make_document
):
    """The failure this pairs with: being accepted at the boundary and dying two
    stages later with "no digitizer", which is what images used to do."""
    from mrip.ingest import digitize as digitizer
    from mrip.schemas import DocumentClass

    if digitizer.ocr_available():
        pytest.skip("the OCR extra is installed, so this path succeeds instead")

    path = tmp_path / "scan.png"
    path.write_bytes(
        bytes.fromhex(
            "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
            "0000000a49444154789c63000100000500010d0a2db40000000049454e44ae426082"
        )
    )
    reference = blobs.put_file(path)
    document = store.register_document(
        make_document(
            "scan.png",
            document_id="doc_scan",
            content_hash=reference.content_hash,
            size_bytes=reference.size_bytes,
            doc_class=DocumentClass.IMAGE,
            page_count=None,
            blob_key=reference.content_hash,
        )
    )

    run_stage(store, document.document_id, "document.classify")
    with pytest.raises(RuntimeError, match=r"ocr"):
        run_stage(store, document.document_id, "document.digitize")
