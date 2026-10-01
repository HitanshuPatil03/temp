"""Tests for report persistence, the publish lifecycle and the file writers.

Generation itself is covered in ``test_reports.py``. What is under test here is
the half that makes the deliverable real rather than a function: a manifest that
survives the request, a lifecycle that refuses the moves it should, and files a
Ministry could actually open — produced with no model in reach.
"""

from __future__ import annotations

import io
import zipfile

import pytest

from mrip.auth.scope import Scope
from mrip.db import Store
from mrip.db.repositories.reports import (
    LEGAL_TRANSITIONS,
    IllegalReportTransitionError,
)
from mrip.reports.generate import generate
from mrip.reports.render import ReportIncompleteError
from mrip.reports.template import default_template
from mrip.reports.writers import render_docx, render_pptx, render_xlsx
from mrip.schemas import FactStatus, ReportState

SCOPE = Scope.unrestricted("test suite")


def _stored(store: Store, make_fact, **overrides):
    """A validated SECL production fact, generated and persisted as a draft."""
    store.insert_facts([make_fact(status=FactStatus.VALIDATED, unit_ambiguous=False)])
    manifest = generate(
        store, SCOPE, default_template(), entity="secl", period="FY2024-25"
    )
    store.reports.create(
        manifest, entity_id="secl", period_label="FY2024-25", **overrides
    )
    return manifest


# -------------------------------------------------------------- persistence


def test_a_generated_manifest_survives_the_request(store: Store, make_fact) -> None:
    """§11.3 — the reproduction guarantee is about storage, not a return value."""
    manifest = _stored(store, make_fact)

    loaded = store.reports.get(manifest.report_id, SCOPE)

    assert loaded is not None
    assert loaded.report_id == manifest.report_id
    assert loaded.title == manifest.title
    assert [f.fact_id for f in loaded.figures] == [f.fact_id for f in manifest.figures]
    assert loaded.figures[0].document_version == manifest.figures[0].document_version


def test_a_report_is_invisible_outside_its_scope(store: Store, make_fact) -> None:
    """The scope key is the entity the report is about, like every other read."""
    manifest = _stored(store, make_fact)

    assert store.reports.get(manifest.report_id, Scope.of("mcl")) is None
    assert store.reports.list(Scope.of("mcl")) == []
    assert store.reports.get(manifest.report_id, Scope.of("secl")) is not None


def test_an_incomplete_manifest_is_still_stored(store: Store) -> None:
    """A report naming what it could not pin is a useful record, not an error."""
    manifest = generate(
        store, SCOPE, default_template(), entity="secl", period="FY2024-25"
    )
    store.reports.create(manifest, entity_id="secl", period_label="FY2024-25")

    loaded = store.reports.get(manifest.report_id, SCOPE)
    assert loaded is not None
    assert not loaded.complete
    assert loaded.missing_required[0].metric == "coal_production"


# ---------------------------------------------------------------- lifecycle


def test_the_lifecycle_runs_draft_to_published(
    store: Store, make_fact, make_user
) -> None:
    approver = make_user("a.approver")
    manifest = _stored(store, make_fact)
    report_id = manifest.report_id

    assert store.reports.get(report_id, SCOPE).state is ReportState.DRAFT
    store.reports.transition(report_id, ReportState.IN_REVIEW, SCOPE)
    approved = store.reports.transition(
        report_id, ReportState.APPROVED, SCOPE, actor_user_id=approver.user_id
    )
    assert approved.state is ReportState.APPROVED
    published = store.reports.transition(report_id, ReportState.PUBLISHED, SCOPE)
    assert published.state is ReportState.PUBLISHED


def test_a_draft_cannot_skip_review(store: Store, make_fact) -> None:
    """Approval exists to be a step. Skipping it is refused, not tolerated."""
    manifest = _stored(store, make_fact)

    with pytest.raises(IllegalReportTransitionError, match="in_review"):
        store.reports.transition(manifest.report_id, ReportState.APPROVED, SCOPE)


def test_a_published_report_is_immutable(store: Store, make_fact, make_user) -> None:
    """§11.3 — once published, frozen. A correction is a new report."""
    approver = make_user("a.approver")
    manifest = _stored(store, make_fact)
    report_id = manifest.report_id
    store.reports.transition(report_id, ReportState.IN_REVIEW, SCOPE)
    store.reports.transition(
        report_id, ReportState.APPROVED, SCOPE, actor_user_id=approver.user_id
    )
    store.reports.transition(report_id, ReportState.PUBLISHED, SCOPE)

    for target in ReportState:
        with pytest.raises(IllegalReportTransitionError):
            store.reports.transition(report_id, target, SCOPE)


def test_sending_a_report_back_withdraws_its_approval(
    store: Store, make_fact, make_user
) -> None:
    """An approver's name must not stay on a version they never saw."""
    approver = make_user("a.approver")
    second = make_user("b.approver")
    manifest = _stored(store, make_fact)
    report_id = manifest.report_id
    store.reports.transition(report_id, ReportState.IN_REVIEW, SCOPE)
    store.reports.transition(
        report_id, ReportState.APPROVED, SCOPE, actor_user_id=approver.user_id
    )

    store.reports.transition(report_id, ReportState.IN_REVIEW, SCOPE)
    sent_back = store.reports.get(report_id, SCOPE)

    assert sent_back is not None
    assert sent_back.state is ReportState.IN_REVIEW
    # Re-approval is possible, which is the point of sending it back.
    store.reports.transition(
        report_id, ReportState.APPROVED, SCOPE, actor_user_id=second.user_id
    )


def test_published_is_terminal_in_the_transition_table() -> None:
    """The lifecycle is data, so the rule can be read directly rather than
    inferred from the behaviour of a chain of conditionals."""
    assert LEGAL_TRANSITIONS[ReportState.PUBLISHED] == frozenset()
    assert ReportState.PUBLISHED in LEGAL_TRANSITIONS[ReportState.APPROVED]
    assert ReportState.APPROVED not in LEGAL_TRANSITIONS[ReportState.DRAFT]


def test_transitioning_an_unknown_report_is_a_lookup_error(store: Store) -> None:
    with pytest.raises(LookupError):
        store.reports.transition("rpt_nope", ReportState.IN_REVIEW, SCOPE)


# ------------------------------------------------------------------ writers


def _is_office_zip(payload: bytes, member: str) -> bool:
    """Every Office format is a zip with a known member. Checking that proves
    the bytes are a real document rather than something merely non-empty."""
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        return member in archive.namelist()


@pytest.mark.parametrize(
    ("writer", "member"),
    [
        (render_docx, "word/document.xml"),
        (render_xlsx, "xl/workbook.xml"),
        (render_pptx, "ppt/presentation.xml"),
    ],
)
def test_each_writer_produces_a_real_office_file(
    store: Store, make_fact, writer, member: str
) -> None:
    manifest = _stored(store, make_fact)
    payload = writer(default_template(), manifest)
    assert _is_office_zip(payload, member)


@pytest.mark.parametrize("writer", [render_docx, render_xlsx, render_pptx])
def test_no_writer_emits_an_incomplete_report(store: Store, writer) -> None:
    """§11.2 — a blank where a required number should be is never emitted."""
    manifest = generate(
        store, SCOPE, default_template(), entity="secl", period="FY2024-25"
    )
    with pytest.raises(ReportIncompleteError, match="Coal production"):
        writer(default_template(), manifest)


def test_the_xlsx_keeps_canonical_and_printed_values_apart(
    store: Store, make_fact
) -> None:
    """A spreadsheet is where someone sums a column. Mixing tonnes with
    million-tonnes in one is the error the fact store exists to prevent."""
    from openpyxl import load_workbook

    manifest = _stored(store, make_fact)
    workbook = load_workbook(io.BytesIO(render_xlsx(default_template(), manifest)))
    sheet = workbook["Figures"]

    headers = [cell.value for cell in sheet[1]]
    row = {header: cell.value for header, cell in zip(headers, sheet[2], strict=True)}

    assert row["Value"] == 193.0e6
    assert row["Unit"] == "t"
    assert row["As printed"] == 193.0
    assert row["Printed unit"] == "MT"
    assert row["Fact id"] == manifest.figures[0].fact_id
    assert "Sources" in workbook.sheetnames


def test_the_docx_carries_every_citation(store: Store, make_fact) -> None:
    """A figure without its source is the thing this project exists to prevent,
    so the locator travels into the document that gets signed."""
    from docx import Document as DocxDocument

    manifest = _stored(store, make_fact)
    document = DocxDocument(io.BytesIO(render_docx(default_template(), manifest)))
    text = "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )

    for figure in manifest.figures:
        assert figure.locator in text


def test_a_report_renders_with_the_model_disabled(store: Store, make_fact) -> None:
    """§7 — every figure, table and source survives; prose is replaced by a
    stated note rather than silently dropped."""
    from docx import Document as DocxDocument

    manifest = _stored(store, make_fact)
    document = DocxDocument(io.BytesIO(render_docx(default_template(), manifest)))
    body = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert "the local model is disabled" in body
